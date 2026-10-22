from __future__ import annotations

import os
import threading
from pathlib import Path
from io import BytesIO
from typing import Any

from dotenv import load_dotenv
from fastapi import FastAPI, BackgroundTasks, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from api.retrieval import answer_question, load_chunks, load_latest_report
from ingestion.run_ingestion import run_ingestion

load_dotenv()
ROOT = Path(__file__).resolve().parents[1]

app = FastAPI(title="STIF.ai SVCE College Assistant", version="0.1.0")

# Permit the local Vite frontend to call this API.
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["Content-Type", "Authorization"],
)

app.mount("/static", StaticFiles(directory=str(ROOT / "frontend")), name="static")

_ingest_lock = threading.Lock()
_ingest_state: dict[str, Any] = {
    "status": "idle",
    "message": "No ingestion job started in this process.",
}


class ChatRequest(BaseModel):
    message: str = Field(min_length=1, max_length=3000)


class ChatResponse(BaseModel):
    answer: str
    sources: list[dict[str, str]] = []
    mode: str = "retrieval"


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".png", ".jpg", ".jpeg", ".webp"}

def _extract_uploaded_text(filename: str, content: bytes) -> tuple[str, str]:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Unsupported file type.")
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return content.decode("utf-8-sig", errors="replace"), suffix[1:]
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages), "pdf"
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read this PDF. It may be invalid or password-protected.") from exc
    if suffix == ".docx":
        try:
            from docx import Document
            doc = Document(BytesIO(content))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows: parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            return "\n".join(parts), "docx"
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read this DOCX document.") from exc
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            from PIL import Image
            import pytesseract
            return pytesseract.image_to_string(Image.open(BytesIO(content)).convert("RGB")), "image-ocr"
        except ImportError as exc:
            raise HTTPException(status_code=503, detail="Image OCR needs Pillow, pytesseract, and the Tesseract OCR application installed on the backend computer.") from exc
        except Exception as exc:
            raise HTTPException(status_code=422, detail="Could not read text from this image. Try a clearer image.") from exc
    raise HTTPException(status_code=415, detail="Unsupported file type.")


@app.get("/", response_class=HTMLResponse)
def home():
    return (ROOT / "frontend" / "index.html").read_text(encoding="utf-8")


@app.get("/health")
def health():
    chunks = load_chunks()
    return {
        "status": "ok",
        "project": "STIF.ai SVCE College Assistant",
        "college": os.getenv(
            "COLLEGE_NAME",
            "Sri Venkateswara College of Engineering, Tirupati",
        ),
        "chunks_loaded": len(chunks),
        "llm_configured": bool(
            os.getenv("LLM_API_KEY") and os.getenv("LLM_MODEL")
        ),
    }


@app.post("/api/v1/chat", response_model=ChatResponse)
def chat(request: ChatRequest):
    result = answer_question(request.message.strip())
    return ChatResponse(**result)


@app.post("/api/v1/files/extract")
async def extract_uploaded_file(file: UploadFile = File(...)):
    filename = Path(file.filename or "upload").name
    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if not content: raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES: raise HTTPException(status_code=413, detail="Maximum upload size is 10 MB.")
    text, file_type = _extract_uploaded_text(filename, content)
    return {"filename": filename, "file_type": file_type, "text": text[:50000], "truncated": len(text) > 50000}


@app.get("/api/v1/ingest/status")
def ingestion_status():
    return {"job": dict(_ingest_state), "last_report": load_latest_report()}


def _run_ingest_task(start_url: str, max_pages: int):
    global _ingest_state
    try:
        _ingest_state.update(status="running", message="Ingestion is running.")
        report = run_ingestion(start_url=start_url, max_pages=max_pages)
        _ingest_state.update(
            status="completed",
            message="Ingestion finished.",
            report=report,
        )
    except Exception as exc:
        _ingest_state.update(status="failed", message=str(exc))


@app.post("/api/v1/ingest/run")
def start_ingestion(
    background_tasks: BackgroundTasks,
    max_pages: int | None = None,
):
    global _ingest_state

    if _ingest_state.get("status") in {"queued", "running"}:
        return {
            "status": "running",
            "message": "An ingestion job is already running.",
        }

    if not _ingest_lock.acquire(blocking=False):
        return {
            "status": "running",
            "message": "An ingestion job is already running.",
        }

    def wrapped():
        try:
            _run_ingest_task(
                os.getenv("START_URL", "https://svce.edu.in/"),
                max_pages or int(os.getenv("MAX_PAGES", "500")),
            )
        finally:
            _ingest_lock.release()

    _ingest_state = {"status": "queued", "message": "Ingestion job queued."}
    background_tasks.add_task(wrapped)
    return {
        "status": "queued",
        "message": "Ingestion job queued. Check /api/v1/ingest/status.",
    }
